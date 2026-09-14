"""班级归属、删除权限、学生名单导入导出。

三条主线：
1. 公共资源（题库、练习文本、打字成绩）没有删除权就删不掉；
   更新日志不在其中 —— 那整个模块只有管理员碰得到，见 test_community；
2. 老师只能给自己名下的班发考试，管理员发的全体学生都收得到；
3. 学生名单能用 Excel/CSV 导进导出，三列：学号、姓名、班级。
"""

import csv
import io
import itertools

import pytest
from openpyxl import Workbook, load_workbook

XLSX_MAGIC = b"PK\x03\x04"

# 本模块建过的题目 id。题库是全局共享的，留在里面会被别的模块抽进卷子 ——
# 而这里的判断题答案写法和那边的约定不一样，会把人家的满分用例判成 75 分。
_made_questions: list[int] = []


@pytest.fixture(scope="module", autouse=True)
def _cleanup_questions(client, auth):
    """跑完把本模块建的题清掉，不给后面的模块留垃圾。"""
    yield
    for qid in _made_questions:
        client.delete(f"/api/questions/{qid}", headers=auth)


def _cid(client, auth, grade, name):
    for c in client.get("/api/classes", headers=auth).json():
        if c["grade"] == grade and c["name"] == name:
            return c["id"]
    return client.post(
        "/api/classes", json={"grade": grade, "name": name}, headers=auth
    ).json()["id"]


_seq = itertools.count(1)


def _a_question(client, auth):
    """造一道一次性的题。题干必须每次不同 —— 题库按题干哈希查重，重复会 409。"""
    res = client.post(
        "/api/questions",
        json={
            "type": "判断题",
            "stem": f"删除权限测试用题 {next(_seq)}：信息技术课很有用。",
            "answer": "对",
            "scope": "信息基础与信息技术",
        },
        headers=auth,
    )
    assert res.status_code == 200, res.text
    qid = res.json()["id"]
    _made_questions.append(qid)
    return qid


# ================================================================ 删除权限
def test_teacher_cannot_delete_shared_resources(client, auth, teacher_auth):
    """默认状态下老师删不动公共资源，而且要说清楚该找谁。"""
    qid = _a_question(client, auth)
    text_id = client.post(
        "/api/typing/texts",
        json={"mode": "chinese", "difficulty": "简单", "content": "删除权限测试用的一段文本。"},
        headers=auth,
    ).json()["id"]
    for url in [
        f"/api/questions/{qid}",
        f"/api/typing/texts/{text_id}",
        "/api/typing/records",
    ]:
        res = client.delete(url, headers=teacher_auth)
        assert res.status_code == 403, f"{url} 竟然让没权限的老师删掉了"
        detail = res.json()["detail"]
        assert "删除权" in detail and "管理员" in detail, f"{url} 的提示没说清该怎么办：{detail}"


def test_teacher_can_still_create_and_edit(client, teacher_auth):
    """挡的只是删除。老师照常建题、改题，不然题库就没法维护了。"""
    res = client.post(
        "/api/questions",
        json={
            "type": "判断题",
            "stem": "老师仍然可以新增题目，这条应该建得成。",
            "answer": "对",
            "scope": "信息基础与信息技术",
        },
        headers=teacher_auth,
    )
    assert res.status_code == 200, res.text
    qid = res.json()["id"]
    _made_questions.append(qid)

    upd = client.put(
        f"/api/questions/{qid}",
        json={"stem": "老师仍然可以修改题目，这条应该改得成。"},
        headers=teacher_auth,
    )
    assert upd.status_code == 200, upd.text


def test_admin_grants_delete_permission(client, auth, teacher2_auth):
    """管理员开了删除权，同一位老师立刻删得动。"""
    me = client.get("/api/auth/me", headers=teacher2_auth).json()
    assert me["can_delete"] is False, "默认不该有删除权"

    qid = _a_question(client, auth)
    assert client.delete(f"/api/questions/{qid}", headers=teacher2_auth).status_code == 403

    client.patch(f"/api/auth/users/{me['id']}", json={"can_delete": True}, headers=auth)
    assert client.get("/api/auth/me", headers=teacher2_auth).json()["can_delete"] is True
    assert client.delete(f"/api/questions/{qid}", headers=teacher2_auth).status_code == 200

    # 收回权限后又删不动了
    client.patch(f"/api/auth/users/{me['id']}", json={"can_delete": False}, headers=auth)
    qid2 = _a_question(client, auth)
    assert client.delete(f"/api/questions/{qid2}", headers=teacher2_auth).status_code == 403


def test_admin_never_needs_the_flag(client, auth):
    qid = _a_question(client, auth)
    assert client.get("/api/auth/me", headers=auth).json()["can_delete"] is False
    # 管理员不看这个字段，照样删得动
    assert client.delete(f"/api/questions/{qid}", headers=auth).status_code == 200


def test_bulk_text_delete_needs_permission_but_toggle_does_not(client, auth, teacher_auth):
    tid = client.post(
        "/api/typing/texts",
        json={"mode": "chinese", "difficulty": "简单", "content": "批量删除权限测试用文本。"},
        headers=auth,
    ).json()["id"]

    # 停用不需要权限
    res = client.post(
        "/api/typing/texts/bulk", json={"ids": [tid], "action": "disable"}, headers=teacher_auth
    )
    assert res.status_code == 200, res.text

    # 删除需要
    res = client.post(
        "/api/typing/texts/bulk", json={"ids": [tid], "action": "delete"}, headers=teacher_auth
    )
    assert res.status_code == 403
    assert "删除权" in res.json()["detail"]


# ================================================================ 班级归属
def test_teacher_without_class_cannot_publish(client, auth):
    """名下没班的老师发不出考试，提示要告诉他去找管理员。"""
    client.post(
        "/api/auth/users",
        json={"username": "noclass_t", "password": "teacher123", "name": "无班老师"},
        headers=auth,
    )
    hdr = {
        "Authorization": "Bearer "
        + client.post(
            "/api/auth/login", data={"username": "noclass_t", "password": "teacher123"}
        ).json()["access_token"]
    }

    pid = client.post(
        "/api/papers/generate",
        json={"title": "无班老师的卷子", "by_sections": False, "counts": {"判断题": 1}, "save": True},
        headers=hdr,
    ).json()["paper_id"]

    res = client.post("/api/exams", json={"paper_id": pid}, headers=hdr)
    assert res.status_code == 400
    assert "没有班级" in res.json()["detail"] and "管理员" in res.json()["detail"]


def test_teacher_cannot_target_someone_elses_class(client, auth, teacher_auth):
    other = _cid(client, auth, "九年级", "8班")
    pid = client.post(
        "/api/papers/generate",
        json={"title": "想发给别人班的卷子", "by_sections": False, "counts": {"判断题": 1}, "save": True},
        headers=teacher_auth,
    ).json()["paper_id"]

    res = client.post(
        "/api/exams",
        json={"paper_id": pid, "target_class_ids": [other]},
        headers=teacher_auth,
    )
    assert res.status_code == 403
    assert "不是你带的班" in res.json()["detail"]


def test_teacher_default_targets_own_classes(client, auth, teacher_auth):
    """老师不选班时，默认发给自己名下的全部班，而不是全校。"""
    pid = client.post(
        "/api/papers/generate",
        json={"title": "默认发给自己班", "by_sections": False, "counts": {"判断题": 1}, "save": True},
        headers=teacher_auth,
    ).json()["paper_id"]
    exam = client.post("/api/exams", json={"paper_id": pid}, headers=teacher_auth).json()

    assert exam["target_classes"], "老师发的考试不能是「全体学生」"
    assert "七年级1班" in exam["target_classes"]


def test_admin_exam_reaches_everyone(client, auth, teacher_auth):
    """管理员不选班 = 全体学生，哪个班的学生都看得到。"""
    pid = client.post(
        "/api/papers/generate",
        json={"title": "管理员发给全校", "by_sections": False, "counts": {"判断题": 1}, "save": True},
        headers=auth,
    ).json()["paper_id"]
    exam = client.post("/api/exams", json={"paper_id": pid}, headers=auth).json()
    assert exam["target_classes"] == "", "管理员不选班就该是全体"

    # 一个谁的班都不沾的学生也要收得到
    cid = _cid(client, auth, "六年级", "1班")
    client.post(
        "/api/students",
        json={"student_no": "26060001", "name": "六年级学生", "class_id": cid},
        headers=auth,
    )
    tok = client.post(
        "/api/student/login", params={"student_no": "26060001", "password": "26060001"}
    ).json()["access_token"]
    seen = client.get(
        "/api/student/exams", headers={"Authorization": f"Bearer {tok}"}
    ).json()
    assert exam["id"] in {e["id"] for e in seen}


def test_teacher_exam_only_reaches_own_class(client, auth, teacher_auth):
    pid = client.post(
        "/api/papers/generate",
        json={"title": "只给甲老师的班", "by_sections": False, "counts": {"判断题": 1}, "save": True},
        headers=teacher_auth,
    ).json()["paper_id"]
    exam = client.post("/api/exams", json={"paper_id": pid}, headers=teacher_auth).json()

    def sees(no, cid):
        client.post(
            "/api/students",
            json={"student_no": no, "name": "某学生", "class_id": cid},
            headers=auth,
        )
        tok = client.post(
            "/api/student/login", params={"student_no": no, "password": no}
        ).json()["access_token"]
        rows = client.get(
            "/api/student/exams", headers={"Authorization": f"Bearer {tok}"}
        ).json()
        return exam["id"] in {e["id"] for e in rows}

    assert sees("26070101", _cid(client, auth, "七年级", "1班")), "自己班的学生该收到"
    assert not sees("26070201", _cid(client, auth, "七年级", "2班")), "别人班的学生不该收到"


def test_only_admin_manages_classes(client, teacher_auth):
    assert client.post(
        "/api/classes", json={"grade": "十年级", "name": "1班"}, headers=teacher_auth
    ).status_code == 403
    # 但老师看得到列表 —— 发考试要选班
    assert client.get("/api/classes", headers=teacher_auth).status_code == 200


def test_class_rename_updates_students(client, auth):
    cid = _cid(client, auth, "五年级", "3班")
    client.post(
        "/api/students",
        json={"student_no": "26050301", "name": "改名测试", "class_id": cid},
        headers=auth,
    )
    client.patch(f"/api/classes/{cid}", json={"name": "4班"}, headers=auth)

    rows = client.get("/api/students", params={"class_id": cid}, headers=auth).json()
    assert rows[0]["student_class"] == "五年级4班", "班级改名后学生的班级名要跟着改"


def test_class_with_students_needs_force_to_delete(client, auth):
    cid = _cid(client, auth, "四年级", "1班")
    client.post(
        "/api/students",
        json={"student_no": "26040101", "name": "在册学生", "class_id": cid},
        headers=auth,
    )

    res = client.delete(f"/api/classes/{cid}", headers=auth)
    assert res.status_code == 409
    assert "还有 1 名学生" in res.json()["detail"]

    res = client.delete(f"/api/classes/{cid}", params={"force": True}, headers=auth)
    assert res.status_code == 200

    # 学生还在，只是变成未分班
    rows = client.get("/api/students", params={"keyword": "26040101"}, headers=auth).json()
    assert len(rows) == 1 and rows[0]["class_id"] is None


def test_batch_create_classes(client, auth):
    res = client.post(
        "/api/classes/batch", json={"grade": "三年级", "start": 1, "end": 5}, headers=auth
    )
    assert res.json()["added"] == 5
    # 再来一次全是重复
    again = client.post(
        "/api/classes/batch", json={"grade": "三年级", "start": 1, "end": 5}, headers=auth
    )
    assert again.json() == {"added": 0, "skipped": 5}


# ================================================================ 导入导出
def test_template_has_three_columns_and_lists_classes(client, auth):
    """模板三列：学号、姓名、班级，并把现有班级原样列出来供照抄。"""
    _cid(client, auth, "七年级", "9班")
    res = client.get("/api/students/template.xlsx", headers=auth)
    assert res.status_code == 200
    assert res.content.startswith(XLSX_MAGIC)

    wb = load_workbook(io.BytesIO(res.content))
    assert [c.value for c in wb["学生名单"][1]] == ["学号", "姓名", "班级"]

    # 「可用班级」那一页要能查到刚建的班，老师复制粘贴就不会写错
    names = {row[0] for row in wb["可用班级"].iter_rows(min_row=2, values_only=True)}
    assert "七年级9班" in names


def _xlsx(rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_xlsx(client, auth):
    cid = _cid(client, auth, "二年级", "1班")
    data = _xlsx([
        ["学号", "姓名"],          # 表头要被跳过
        ["26020101", "甲同学"],
        ["26020102", "乙同学"],
        ["", "缺学号"],            # 报错但不影响别的行
        ["26020103", "丙同学"],
    ])
    res = client.post(
        "/api/students/import",
        params={"class_id": cid},
        files={"file": ("名单.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth,
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["added"] == 3
    assert body["error_count"] == 1
    assert body["student_class"] == "二年级1班"


def test_import_keeps_student_no_as_written(client, auth):
    """Excel 把学号存成数字时，读出来是 26020201.0，不能带小数点。"""
    cid = _cid(client, auth, "二年级", "2班")
    data = _xlsx([["学号", "姓名"], [26020201, "数字学号"]])
    client.post(
        "/api/students/import",
        params={"class_id": cid},
        files={"file": ("名单.xlsx", data, "application/octet-stream")},
        headers=auth,
    )
    rows = client.get("/api/students", params={"class_id": cid}, headers=auth).json()
    assert [r["student_no"] for r in rows] == ["26020201"]


@pytest.mark.parametrize("encoding", ["utf-8-sig", "gbk"])
def test_import_csv(client, auth, encoding):
    cid = _cid(client, auth, "一年级", "1班" if encoding == "gbk" else "2班")
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["学号", "姓名"])
    w.writerow([f"2601{encoding[:2]}01", "钱同学"])
    raw = buf.getvalue().encode(encoding)

    res = client.post(
        "/api/students/import",
        params={"class_id": cid},
        files={"file": ("名单.csv", raw, "text/csv")},
        headers=auth,
    )
    assert res.status_code == 200, res.text
    assert res.json()["added"] == 1


def test_import_rejects_other_formats(client, auth):
    res = client.post(
        "/api/students/import",
        files={"file": ("名单.pdf", b"%PDF-1.4", "application/pdf")},
        headers=auth,
    )
    assert res.status_code == 400
    assert ".xlsx" in res.json()["detail"]


def test_export_round_trips(client, auth):
    """导出的表前两列和模板一致，改完能直接再导回来。"""
    cid = _cid(client, auth, "二年级", "3班")
    client.post(
        "/api/students/import",
        params={"class_id": cid},
        files={"file": ("名单.xlsx", _xlsx([["学号", "姓名"], ["26020301", "导出测试"]]), "application/octet-stream")},
        headers=auth,
    )

    res = client.get("/api/students/export.xlsx", params={"class_id": cid}, headers=auth)
    assert res.status_code == 200
    assert res.content.startswith(XLSX_MAGIC)

    ws = load_workbook(io.BytesIO(res.content)).active
    assert [c.value for c in ws[1]][:2] == ["学号", "姓名"]
    assert [c.value for c in ws[2]][:3] == ["26020301", "导出测试", "二年级3班"]

    # 把导出的表原样再导一次，应该全部识别为重复而不是报错
    again = client.post(
        "/api/students/import",
        params={"class_id": cid},
        files={"file": ("回传.xlsx", res.content, "application/octet-stream")},
        headers=auth,
    )
    assert again.json()["added"] == 0
    assert again.json()["skipped"] == 1


def test_import_reads_class_from_each_row(client, auth):
    """表里逐行写班级，学生按各自那一行的班级入库。"""
    _cid(client, auth, "五年级", "1班")
    _cid(client, auth, "五年级", "2班")
    data = _xlsx([
        ["学号", "姓名", "班级"],
        ["26050101", "一班甲", "五年级1班"],
        ["26050201", "二班乙", "五年级2班"],
    ])
    res = client.post(
        "/api/students/import",
        files={"file": ("名单.xlsx", data, "application/octet-stream")},
        headers=auth,
    )
    assert res.status_code == 200, res.text
    assert res.json()["added"] == 2
    assert res.json()["classes"] == ["五年级1班", "五年级2班"]

    rows = {
        r["student_no"]: r["student_class"]
        for r in client.get("/api/students", params={"keyword": "2605"}, headers=auth).json()
    }
    assert rows["26050101"] == "五年级1班"
    assert rows["26050201"] == "五年级2班"


def test_import_reports_bad_class_per_row(client, auth):
    """班级写错只跳过那一行，其余照常导入，并指出错在第几行、错成了什么。"""
    _cid(client, auth, "六年级", "3班")
    data = _xlsx([
        ["学号", "姓名", "班级"],
        ["26060301", "写对了", "六年级3班"],
        ["26060302", "写错了", "六(3)班"],      # 系统里没有这个写法
        ["26060303", "也对", "六年级3班"],
    ])
    res = client.post(
        "/api/students/import",
        files={"file": ("名单.xlsx", data, "application/octet-stream")},
        headers=auth,
    )
    body = res.json()
    assert body["added"] == 2, "写错的那一行不该拖垮其余两行"
    assert body["error_count"] == 1

    err = body["errors"][0]
    assert "第 3 行" in err and "六(3)班" in err
    # 每行只说"哪一行、错在哪"，别把怎么改也跟在每一行后面
    assert "完全一致" not in err

    # "该怎么改"整体只给一句，放在 hint 里
    assert "完全一致" in body["hint"] and "六(3)班" in body["hint"]

    # 写错的那个学生确实没被建出来
    got = client.get("/api/students", params={"keyword": "26060302"}, headers=auth).json()
    assert got == []


def test_import_does_not_repeat_the_same_advice_on_every_row(client, auth):
    """十几行都写错班级时，"该怎么改"那句话只能出现一次。

    以前是每行都跟一遍"班级名要和「班级」页面完全一致（如 …）"，
    错十行就刷十遍，真正有用的行号反而被淹了。
    """
    _cid(client, auth, "五年级", "7班")
    rows = [["学号", "姓名", "班级"]]
    rows += [[f"2605070{i}", f"学生{i}", "五(7)班"] for i in range(1, 7)]
    res = client.post(
        "/api/students/import",
        files={"file": ("名单.xlsx", _xlsx(rows), "application/octet-stream")},
        headers=auth,
    )
    body = res.json()
    assert body["error_count"] == 6 and body["added"] == 0

    # 六行错误，六条各自带行号，但没有一条重复那段指导语
    assert len(body["errors"]) == 6
    assert all("第 " in e and "五(7)班" in e for e in body["errors"])
    assert sum("完全一致" in e for e in body["errors"]) == 0

    # 错的班名只列一次，不是列六遍
    assert body["hint"].count("五(7)班") == 1


def test_import_falls_back_to_selected_class(client, auth):
    """班级列留空时，归到导入时选的那个班。"""
    cid = _cid(client, auth, "四年级", "2班")
    data = _xlsx([["学号", "姓名", "班级"], ["26040201", "没写班级", ""]])
    res = client.post(
        "/api/students/import",
        params={"class_id": cid},
        files={"file": ("名单.xlsx", data, "application/octet-stream")},
        headers=auth,
    )
    assert res.json()["added"] == 1
    rows = client.get("/api/students", params={"class_id": cid}, headers=auth).json()
    assert [r["student_no"] for r in rows] == ["26040201"]


def test_import_tolerates_spaces_in_class_name(client, auth):
    """「七年级 1班」这种手滑多打的空格要认出来；真写错的仍然报错。"""
    _cid(client, auth, "七年级", "1班")
    data = _xlsx([["学号", "姓名", "班级"], ["26070199", "多了空格", " 七年级 1班 "]])
    res = client.post(
        "/api/students/import",
        files={"file": ("名单.xlsx", data, "application/octet-stream")},
        headers=auth,
    )
    assert res.json()["added"] == 1
    rows = client.get("/api/students", params={"keyword": "26070199"}, headers=auth).json()
    assert rows[0]["student_class"] == "七年级1班"


def test_export_round_trips_with_class(client, auth):
    """导出的表前三列和模板同构，原样再导回来应该全部识别为重复。"""
    _cid(client, auth, "三年级", "9班")
    client.post(
        "/api/students/import",
        files={
            "file": (
                "名单.xlsx",
                _xlsx([["学号", "姓名", "班级"], ["26030901", "回流测试", "三年级9班"]]),
                "application/octet-stream",
            )
        },
        headers=auth,
    )

    out = client.get(
        "/api/students/export.xlsx", params={"keyword": "26030901"}, headers=auth
    )
    ws = load_workbook(io.BytesIO(out.content)).active
    assert [c.value for c in ws[1]][:3] == ["学号", "姓名", "班级"]
    assert [c.value for c in ws[2]][:3] == ["26030901", "回流测试", "三年级9班"]

    again = client.post(
        "/api/students/import",
        files={"file": ("回传.xlsx", out.content, "application/octet-stream")},
        headers=auth,
    )
    body = again.json()
    assert body["added"] == 0
    assert body["skipped"] == 1
    assert body["error_count"] == 0
