"""打字训练：学生免登录练习与交卷、教师端统计与导出。"""

import io

from openpyxl import load_workbook

XLSX_MAGIC = b"PK\x03\x04"


def _submit(client, **kw):
    body = {
        "student_name": "张三",
        "student_class": "七(3)班",
        "module": "英文",
        "difficulty": "简单",
        "typed_chars": 100,
        "correct_chars": 95,
        "duration": 60,
    }
    body.update(kw)
    return client.post("/api/typing/records", json=body)


# ---------- 学生端：不需要登录 ----------
def test_config_is_public(client):
    res = client.get("/api/typing/config")
    assert res.status_code == 200
    body = res.json()
    assert body["difficulties"] == ["简单", "中等", "困难"]
    # 「不限时」已经取消，档位必须全是正数（不限时交给自由打字模块）
    assert body["time_limits"] and all(t > 0 for t in body["time_limits"])
    assert body["default_limit"] in body["time_limits"]


def test_passage_is_public_and_random(client):
    a = client.get("/api/typing/passage", params={"mode": "chinese", "difficulty": "简单"})
    assert a.status_code == 200
    text = a.json()["text"]
    assert len(text) > 30  # 是拼起来的长文，不是单句

    # 随机洗牌，连着取多次不该次次相同
    seen = {
        client.get(
            "/api/typing/passage", params={"mode": "chinese", "difficulty": "简单"}
        ).json()["text"]
        for _ in range(8)
    }
    assert len(seen) > 1


def test_passage_unknown_difficulty_is_404(client):
    res = client.get("/api/typing/passage", params={"mode": "english", "difficulty": "地狱"})
    assert res.status_code == 404


def test_submit_without_login(client):
    res = _submit(client)
    assert res.status_code == 200
    body = res.json()
    assert body["accuracy"] == 95          # 95/100
    assert body["speed"] == 95             # 95 字 / 1 分钟
    assert body["stars"] == 3              # 简单难度 68+12=80，95 超过


def test_name_and_class_required(client):
    assert _submit(client, student_name=" ").status_code == 400
    assert _submit(client, student_class="").status_code == 400


def test_unknown_module_rejected(client):
    assert _submit(client, module="五笔").status_code == 400


def test_keyboard_module_has_no_speed(client):
    """键盘模块统计的是按键次数，算字/分没有意义。"""
    body = _submit(client, module="键盘", typed_chars=60, correct_chars=54).json()
    assert body["speed"] == 0
    assert body["accuracy"] == 90


def test_correct_cannot_exceed_typed(client):
    """伪造一个 correct > typed 的上报，正确率不能超过 100%。"""
    body = _submit(client, typed_chars=50, correct_chars=999).json()
    assert body["accuracy"] == 100


def test_zero_typed_is_zero_accuracy(client):
    """一个字没敲不能算满分。"""
    body = _submit(client, typed_chars=0, correct_chars=0).json()
    assert body["accuracy"] == 0
    assert body["stars"] == 0


def test_stars_follow_difficulty(client):
    """同样 82% 的正确率，简单难度给三星，困难难度只到两星以下。"""
    easy = _submit(client, difficulty="简单", typed_chars=100, correct_chars=82).json()
    hard = _submit(client, difficulty="困难", typed_chars=100, correct_chars=82).json()
    assert easy["stars"] > hard["stars"]


# ---------- 教师端：要登录 ----------
def test_teacher_endpoints_need_login(client):
    assert client.get("/api/typing/records").status_code == 401
    assert client.get("/api/typing/stats").status_code == 401
    assert client.get("/api/typing/export.xlsx").status_code == 401


def test_records_list_and_filter(client, auth):
    _submit(client, student_name="李四", student_class="八(1)班", module="中文")
    rows = client.get("/api/typing/records", headers=auth).json()
    assert len(rows) >= 2

    only = client.get(
        "/api/typing/records", params={"student_class": "八(1)班"}, headers=auth
    ).json()
    assert only and all(r["student_class"] == "八(1)班" for r in only)

    by_name = client.get(
        "/api/typing/records", params={"keyword": "李"}, headers=auth
    ).json()
    assert by_name and all("李" in r["student_name"] for r in by_name)


def test_records_sorted_by_speed(client, auth):
    rows = client.get("/api/typing/records", params={"order": "speed"}, headers=auth).json()
    speeds = [r["speed"] for r in rows]
    assert speeds == sorted(speeds, reverse=True)


def test_stats(client, auth):
    st = client.get("/api/typing/stats", headers=auth).json()
    assert st["total"] >= 2
    assert st["students"] >= 2
    assert 0 <= st["avg_accuracy"] <= 100
    assert sum(st["accuracy_buckets"].values()) == st["total"]
    assert st["by_class"] and "avg_accuracy" in st["by_class"][0]
    # 键盘记录速度为 0，不能把平均速度拉下去
    assert st["avg_speed"] > 0


def test_classes_list(client, auth):
    got = client.get("/api/typing/classes", headers=auth).json()
    assert "七(3)班" in got and "八(1)班" in got


def test_export_xlsx(client, auth):
    res = client.get("/api/typing/export.xlsx", headers=auth)
    assert res.status_code == 200
    assert res.content.startswith(XLSX_MAGIC)
    wb = load_workbook(io.BytesIO(res.content), read_only=True)
    rows = [list(r) for r in wb["打字成绩"].iter_rows(values_only=True)]
    assert rows[0][:4] == ["姓名", "班级", "模块", "难度"]
    assert len(rows) >= 3


def test_delete_and_clear(client, auth):
    rid = client.get("/api/typing/records", headers=auth).json()[0]["id"]
    assert client.delete(f"/api/typing/records/{rid}", headers=auth).status_code == 200
    assert client.delete(f"/api/typing/records/{rid}", headers=auth).status_code == 404

    res = client.delete(
        "/api/typing/records", params={"student_class": "八(1)班"}, headers=auth
    )
    assert res.status_code == 200 and res.json()["deleted"] >= 1
    left = client.get(
        "/api/typing/records", params={"student_class": "八(1)班"}, headers=auth
    ).json()
    assert left == []
